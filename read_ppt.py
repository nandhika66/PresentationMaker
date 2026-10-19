# STEP 1 - Extract text from every PPTX in input_ppts/
# and save one structured record per slide to output/slide_records.json

import json

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from config import INPUT_DIR, RECORDS_FILE


def get_slide_title(slide):
    """Find the slide title, with a text fallback."""

    # First, look for PowerPoint's title placeholder
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        ):
            title = shape.text.strip()
            if title:
                return title

    # Fall back to the first non-empty text shape
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text.strip()
            if text:
                return text.splitlines()[0][:150]

    return "Untitled slide"


def get_presentation_title(presentation):
    """Extract a title from the first slide of a presentation."""

    if not len(presentation.slides):
        return "Untitled presentation"

    first_slide = presentation.slides[0]
    first_slide_text = [
        shape.text.strip()
        for shape in first_slide.shapes
        if shape.has_text_frame and shape.text.strip()
    ]

    deck_title = " ".join(first_slide_text).strip()
    return deck_title or "Untitled presentation"


def iter_shapes(shapes):
    """Yield every shape, including shapes inside groups."""

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def extract_slide(slide, source_file, slide_number, deck_title):
    """Convert one PowerPoint slide into a structured record."""

    title = get_slide_title(slide)

    text_elements = []
    tables = []
    images = []

    for shape in iter_shapes(slide.shapes):

        # Editable text
        if shape.has_text_frame:
            text = shape.text.strip()
            if text:
                text_elements.append(text)

        # Tables as rows and cells
        if getattr(shape, "has_table", False) and shape.has_table:
            tables.append([
                [cell.text.strip() for cell in row.cells]
                for row in shape.table.rows
            ])

        # Basic picture references (no OCR for now)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images.append({
                "shape_name": shape.name,
                "width": shape.width,
                "height": shape.height,
            })

    # Speaker notes often hold useful text too
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()

    table_text = [
        cell
        for table in tables
        for row in table
        for cell in row
        if cell
    ]

    # Same search text as before (notes are stored separately)
    search_text = "\n".join([title, *text_elements, *table_text])

    return {
        "source_file": source_file,
        "presentation_title": deck_title,
        "slide_number": slide_number,
        "slide_title": title,
        "layout_name": slide.slide_layout.name,
        "text_elements": text_elements,
        "tables": tables,
        "images": images,
        "notes": notes,
        "search_text": search_text,
    }


def main():
    INPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Find PowerPoint files, excluding temporary lock files
    ppt_files = sorted(
        file for file in INPUT_DIR.glob("*.pptx")
        if not file.name.startswith("~$")
    )

    print(f"Found {len(ppt_files)} PowerPoint files in {INPUT_DIR}")

    if not ppt_files:
        print("Put your .pptx files in the input_ppts folder and run again.")
        return

    all_slides = []

    for ppt_path in ppt_files:
        print(f"\nProcessing: {ppt_path.name}")

        try:
            presentation = Presentation(str(ppt_path))
        except Exception as error:
            print(f"  Skipped (could not open): {error}")
            continue

        deck_title = get_presentation_title(presentation)
        print(f"  Presentation title: {deck_title[:80]}")

        for number, slide in enumerate(presentation.slides, start=1):
            all_slides.append(
                extract_slide(slide, ppt_path.name, number, deck_title)
            )

    with RECORDS_FILE.open("w", encoding="utf-8") as file:
        json.dump(all_slides, file, indent=2, ensure_ascii=False)

    print(f"\nProcessed {len(ppt_files)} presentation(s).")
    print(f"Created {len(all_slides)} slide records.")
    print(f"Output: {RECORDS_FILE}")
    print("\nNext step: python build_embeddings.py")


if __name__ == "__main__":
    main()
